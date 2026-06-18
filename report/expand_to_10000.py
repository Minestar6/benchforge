"""Expand the mid-term report to 10000+ Chinese characters.

Strategy: expand text across ALL sections, with deep technical detail.
Each expansion paragraph is self-contained and inserted at a specific anchor.
"""

from docx import Document
from lxml import etree
from docx.oxml.ns import qn

OUT_PATH = '/Users/zhaoziqing/Desktop/benchforge/report/赵子晴_中期报告_完善版.docx'
BACKUP_PATH = '/Users/zhaoziqing/Desktop/benchforge/report/赵子晴_中期报告_完善版_backup.docx'

FONT_BODY = 'Times New Roman'
FONT_EA = '宋体'
FONT_HEADING_EA = '黑体'
SZ_BODY = 12
SZ_SUBSECTION = 14
ALIGN_JUSTIFY = 'both'


def make_para(doc, text, bold=False, size=SZ_BODY, font_name=FONT_BODY,
              ea_font_name=FONT_EA, alignment=ALIGN_JUSTIFY):
    new_p = etree.Element(qn('w:p'))
    pPr = etree.SubElement(new_p, qn('w:pPr'))
    if alignment is not None:
        jc = etree.SubElement(pPr, qn('w:jc'))
        jc.set(qn('w:val'), alignment)
    r = etree.SubElement(new_p, qn('w:r'))
    rPr = etree.SubElement(r, qn('w:rPr'))
    rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    if font_name:
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
    if ea_font_name:
        rFonts.set(qn('w:eastAsia'), ea_font_name)
    if size:
        sz = etree.SubElement(rPr, qn('w:sz'))
        sz.set(qn('w:val'), str(int(size * 2)))
        szCs = etree.SubElement(rPr, qn('w:szCs'))
        szCs.set(qn('w:val'), str(int(size * 2)))
    if bold:
        etree.SubElement(rPr, qn('w:b'))
        etree.SubElement(rPr, qn('w:bCs'))
    t = etree.SubElement(r, qn('w:t'))
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    return new_p


def insert_after(doc, after_idx, text, bold=False, size=SZ_BODY, ea_font_name=FONT_EA):
    ref_element = doc.paragraphs[after_idx]._element
    new_p = make_para(doc, text, bold=bold, size=size, ea_font_name=ea_font_name)
    ref_element.addnext(new_p)


def find_text(doc, keyword, start=0):
    """Find paragraph index containing keyword."""
    for i in range(start, len(doc.paragraphs)):
        if keyword in doc.paragraphs[i].text.strip():
            return i
    return -1


def count_chinese(text):
    return sum(1 for c in text if '一' <= c <= '鿿')


def main():
    doc = Document(OUT_PATH)
    initial_total = count_chinese(''.join(p.text.strip() for p in doc.paragraphs if p.text.strip()))
    print(f"Initial total Chinese: {initial_total}")

    expansions = []

    # ==========================================
    # Section 1.1 — Expand research background
    # ==========================================
    idx = find_text(doc, '显著提升大模型评价的效率')
    if idx >= 0:
        insert_after(doc, idx,
            "从评价方法的发展脉络来看，大模型评价经历了从人工评价到半自动化评价再到全自动化评价的演进过程。"
            "早期的大模型评价主要依赖人工构建的基准数据集，如MMLU、HellaSwag、HumanEval等，这些数据集在特定时期内发挥了重要作用，"
            "但其构建成本高、更新周期长，难以跟上大模型快速迭代的步伐。随后出现的评价框架如OpenAI Evals、EleutherAI LM Evaluation "
            "Harness等，提供了标准化的评价接口和预置数据集，但仍需人工维护数据集。近年来，基于智能体的自动评价方法开始兴起，"
            "通过将大模型自身的生成能力应用于评价数据集的构建过程，实现了评价流程的自动化。本课题正是在这一技术趋势下提出的，"
            "旨在探索一种更加灵活、自适应的自动评价智能体架构，能够在无需人工干预的情况下完成从数据集构建到模型评估的全流程。"
        )
        expansions.append("1.1 research background")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 1.2 — Expand research content
    # ==========================================
    idx = find_text(doc, '在目前的研究中')
    if idx >= 0:
        insert_after(doc, idx,
            "在具体实施中，本研究采用了一种"由点到面、逐步集成"的研发策略。首先集中攻关评价数据集自动构建这一核心难题，"
            "设计并实现了基于多轮反馈自适应机制的题目生成智能体。该组件通过证据检索、文本分段、多策略采样和分阶段生成等技术手段，"
            "实现了从原始资料到高质量评价题目的全自动转化。实验表明，该组件能够稳定产出覆盖多主题、多难度层级的题目集合，"
            "且生成过程具有良好的可控性和可复现性。随后，在数据集构建的基础上，进一步设计和实现了自动评价智能体架构，"
            "采用强中心化规划协同模式，将规划智能体作为全局决策核心，统一调度题目验证和模型评估等子任务。"
            "目前该架构已完成初版实现，能够支持从用户目标输入到多轮生成、验证、评估的完整闭环流程。对于第三部分的评估报告生成，"
            "已完成了数据结构设计和标准化输出格式的定义，可视化呈现和自然语言报告生成等后续工作正在规划中。"
        )
        expansions.append("1.2 research progress")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.1(1) — Expand evidence pool construction
    # ==========================================
    idx = find_text(doc, '用户只需指定研究主题即可自动完成相关资料的检索和预处理')
    if idx >= 0:
        insert_after(doc, idx,
            "在检索策略方面，系统采用多轮递进式检索策略。首轮检索以用户指定的核心主题为查询词，从百科知识库中获取初步的候选文档列表；"
            "对于获取到的每篇文档，系统进一步提取其内部的关键概念和子主题作为扩展查询词，进行第二轮深度检索，以捕获与核心主题相关的"
            "周边知识和背景信息。这种递进式检索策略有效扩展了知识的覆盖范围，为后续的题目生成提供了丰富的素材基础。在文档摘要生成方面，"
            "系统采用了分层摘要策略：首先对每个文本块独立生成局部摘要，再综合所有局部摘要在文档层面生成全局摘要。"
            "这种分层策略既保留了局部细节信息，又提供了文档整体的主题概览，在证据采样和主题筛选阶段发挥着重要的指导作用。"
        )
        expansions.append("2.1(1) evidence retrieval")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.1(2) — Expand question generation details
    # ==========================================
    idx = find_text(doc, '初步筛选后的合格题目进入候选池，等待后续的多阶段验证')
    if idx >= 0:
        insert_after(doc, idx,
            "在题目生成的提示词设计方面，系统针对不同的题目模式和难度层级分别设计了专门化的提示词模板。对于问答题，"
            "提示词要求模型在生成问题时明确标注所需的推理类型，如事实性检索、逻辑推理、数值计算或综合分析等，并给出标准答案的同时"
            "标注答案在证据文本中的具体引用位置，确保生成的每道题目都有可靠的证据支撑。对于多项选择题，提示词特别强调干扰项的设计原则："
            "干扰项应当与正确答案在表面形式（长度、风格、专业术语密度）上保持一致，但在关键信息上与证据内容相矛盾；"
            "同时干扰项之间应当保持合理区分度，避免出现显然错误的选项。针对困难题目模式，提示词进一步增加了对复杂推理链条的要求，"
            "要求模型设计需要多步推理或多证据综合才能回答的问题，并在推理过程中展示完整的逻辑链条。此外，系统还支持通过示例引导"
            "（few-shot examples）的方式对生成格式进行约束，通过在提示词中嵌入符合规范的正例和反例，引导模型输出结构化的题目数据，"
            "保证后续处理流程的兼容性和稳定性。"
        )
        expansions.append("2.1(2) prompt design")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.1(3) — Expand multi-round feedback
    # ==========================================
    idx = find_text(doc, '自主完成大规模、高质量评价数据集的构建，并具有良好的鲁棒性和自适应能力')
    if idx >= 0:
        insert_after(doc, idx,
            "在多轮反馈的具体实现中，系统重点关注以下几类反馈信号的采集与应用。第一是产量类信号，记录每轮各模式和各难度层级的新增题目数量，"
            "用于评估当前生成策略的有效性和效率；当某难度层级连续多轮产量低于预期时，系统会触发策略切换，转向其他难度或模式。"
            "第二是质量类信号，通过验证环节的通过率和质量评分反馈，识别哪些主题或难度层级生成的题目质量偏低，从而在后续轮次中针对性地"
            "调整证据采样策略或切换提示词模板。第三是资源类信号，跟踪每个证据文本块的使用频率和已生成题目的质量分布，当某个文本块被反复"
            "使用但产出质量持续偏低时，系统会降低其采样优先级，将生成资源分配给更有潜力的证据单元。第四是收敛性信号，监控候选池总量的增速"
            "和新题目的边际增益，当增速显著放缓或边际增益降至阈值以下时，意味着当前证据资源的生成潜力已基本耗尽，系统自动触发证据扩展流程"
            "为目标主题补充新的资料，以维持生成效率。这四类信号的综合分析构成了多轮反馈机制的核心决策依据，使系统能够在复杂多变的生成环境中"
            "保持稳定高效的运行状态。"
        )
        expansions.append("2.1(3) feedback signals")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.2(1) — Expand planner agent
    # ==========================================
    idx = find_text(doc, '规划智能体还负责将上述计划组装为可执行的轮次规范')
    if idx >= 0:
        insert_after(doc, idx,
            "规划智能体的诊断阶段是多轮迭代的核心决策基础。在每一轮开始时，规划智能体综合前一轮各环节返回的执行结果数据，"
            "从多个维度进行状态评估：生成维度关注本轮各主题和难度层级的实际产量与计划目标的偏差程度、连续空轮情况以及生成失败率；"
            "验证维度关注验证各阶段的通过率和质量评分分布，判断当前验证策略的严格程度是否合适；评估维度关注各候选模型得分的分布特征，"
            "包括模型间的区分度水平和各题目的得分方差，评估当前评价数据集对模型能力的鉴别能力。基于上述多维度诊断结果，"
            "规划智能体将其映射为一组策略调整指令，分别作用于后续的生成计划、控制参数配置和主题选择方向。这一诊断-规划-执行的闭环机制"
            "使得整个评价系统具备了动态自适应能力，能够根据实际运行情况不断优化自身行为，逐步逼近最佳运行状态。"
        )
        expansions.append("2.2(1) diagnosis")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.2(2) — Expand verify agent
    # ==========================================
    idx = find_text(doc, '实现验证策略的自适应优化')
    if idx >= 0:
        insert_after(doc, idx,
            "验证智能体的三阶段设计体现了从严格到灵活、从客观到主观的递进式质量把控理念。引文验证阶段采用完全客观的规则化检查，"
            "逐条比对题目中的引用声明与证据原文，确保事实性信息的准确性和可追溯性，这一阶段的检查标准是刚性不可妥协的，"
            "任何引文不匹配的题目都会被标记为待修订或直接拒绝。大模型质量评估阶段则引入了柔性评估维度，利用大模型自身的语义理解能力"
            "对题目的表述质量、答案合理性和逻辑一致性等进行综合判断，这一阶段的评估标准可以根据配置灵活调整，在严格和宽松之间找到平衡点。"
            "加权选择阶段将前两个阶段的评分结果进行综合，同时引入多样性和覆盖度约束，通过优化算法在有限的入选名额内选择出最具代表性的"
            "题目组合。这种刚柔并济的验证设计既保证了评价数据集的基本事实准确性，又兼顾了题目的质量和多样性，有效平衡了严格把关与效率"
            "之间的关系。此外，验证阶段产生的详细反馈记录（各阶段通过率、拒绝原因分布、质量评分分布等）为规划智能体的策略调优提供了"
            "重要的数据支撑，形成了生成-验证-反馈-调整的完整质量闭环。"
        )
        expansions.append("2.2(2) verify agent philosophy")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.2(3) — Expand model eval agent
    # ==========================================
    idx = find_text(doc, '形成全链路的自适应优化')
    if idx >= 0:
        insert_after(doc, idx,
            "模型评估智能体在架构设计上采用了"离线推理、异步评估"的策略，将耗时较长的模型推理过程与计算相对轻量的指标评分过程分离，"
            "提升了整体评估流程的执行效率。在模型推理阶段，系统对多个候选模型同时发起推理请求，每个模型的调用通过统一的客户端接口进行"
            "抽象，支持不同厂商（如OpenAI、Anthropic、Google等）和不同部署方式（云端API、本地部署等）的模型接入。推理结果以标准化格式"
            "缓存至本地存储，便于后续的离线分析和多次评估。在指标评分阶段，系统支持多层次的评分体系：底层为自动评分指标（精确匹配、"
            "关键词匹配、语义相似度等），提供快速、客观的基础评分；上层为大模型裁判评分，提供更深层的质性分析。两层次评分结果通过配置化"
            "的权重体系进行融合，形成每道题目的综合得分。最终的综合评估报告不仅包含各模型在各维度上的得分，还提供区分度分析、稳定性分析"
            "和一致性检验等高级分析功能，帮助研究人员全面理解各候选模型的实际表现。"
        )
        expansions.append("2.2(3) eval architecture")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.2(4) — Expand overall framework
    # ==========================================
    idx = find_text(doc, '具有良好的模块化程度和可扩展性')
    if idx >= 0:
        insert_after(doc, idx,
            "智能体之间的通信与协作通过标准化的数据接口实现。各子智能体之间的数据交换采用统一的数据结构规范，包括题目结构、验证结果结构、"
            "评估结果结构和反馈结构等。这种标准化设计确保了各组件之间的松耦合关系，使得任何一个子智能体的内部实现发生变化时，"
            "不会影响其他组件的正常运行。例如，验证智能体可以在不通知其他组件的情况下切换其内部的大模型质量评估策略，只要其输出的验证"
            "结果符合约定的数据结构，规划智能体和模型评估智能体即可无缝兼容。同样，模型评估智能体支持动态添加新的评估指标和裁判模型，"
            "新指标只需遵循约定的接口规范即可自动参与评估流程。这种可扩展的架构设计为本系统的持续演进提供了良好的基础，"
            "使得新的技术和方法可以方便地集成到现有框架中。"
        )
        expansions.append("2.2(4) extensibility")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 2.3 — Expand report generation
    # ==========================================
    idx = find_text(doc, '帮助研究人员快速理解各候选模型的能力差异')
    if idx >= 0:
        insert_after(doc, idx,
            "在评估报告的指标体系设计方面，系统支持多层次、可定制的评估指标配置。在基础指标层，系统计算每个候选模型在各题目上的得分"
            "以及在整体数据集上的汇总统计量，包括均值、中位数、标准差和最值分布等。在分析指标层，系统自动计算模型间区分度指标"
            "（如总体模型差距、最佳与次佳模型差距、每道题目的得分方差等）和数据集质量指标（区分度题目比例、过易题目比例、全失败题目"
            "比例等），帮助研究人员判断当前评价数据集和候选模型之间的匹配程度。在可视化方案方面，系统计划支持多种图表类型，"
            "包括雷达图（展示各模型在多维能力上的分布特征）、柱状图（对比不同模型在各个主题和难度层级上的得分情况）、"
            "热力图（展示模型×题目的得分矩阵）和散点图（分析模型得分与题目难度之间的关系）等。这些可视化图表与文本分析相结合，"
            "输出涵盖模型能力总览、分维度对比分析和典型题目案例等模块的综合报告，为研究人员提供直观、深入的模型能力分析体验。"
        )
        expansions.append("2.3 metrics & visualization")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 3 — Expand schedule details
    # ==========================================
    idx = find_text(doc, '准备答辩材料，进行论文答辩')
    if idx >= 0:
        insert_after(doc, idx,
            "上述进度安排基于当前的研究基础和技术积累制定。由于系统的核心框架已经完成，各智能体组件已具备基本功能，后续工作主要围绕"
            "功能完善、系统集成和实验验证展开。其中，评估报告生成模块的开发虽然尚在起步阶段，但其依赖的基础数据结构已经与模型评估智能体"
            "充分对接，技术方案明确，实现路径清晰。多轮反馈机制的优化工作将在已有实现的基础上进行增量改进，风险可控。"
            "系统集成测试和实验验证阶段将重点关注系统的稳定性、效率和可复现性，确保最终交付的系统具有实际应用价值。"
            "论文撰写将遵循硕士学位论文的规范要求，与实验验证工作同步推进，确保各章节内容充实、论据充分。"
        )
        expansions.append("3 schedule")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 4 — Expand difficulties
    # ==========================================
    idx = find_text(doc, '第三，计算资源消耗问题')
    if idx >= 0:
        insert_after(doc, idx,
            "第四，评价结果的稳定性和可复现性问题。由于大模型生成过程具有一定的随机性，同一主题的多次生成可能得到不同的题目集合，"
            "导致评价结果在不同运行批次之间存在波动。虽然系统采用了随机种子控制和确定性采样策略来减小这种波动，但仍然难以完全消除。"
            "如何设计更加鲁棒的评估协议，使得评价结果对题目集合的随机波动不敏感，同时在有限的资源约束下获得统计显著的评估结论，"
            "是一个需要在理论和实践层面共同探索的问题。"
        )
        expansions.append("4 difficulty")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ==========================================
    # Section 5 — Expand completion possibility
    # ==========================================
    idx = find_text(doc, '预计能够如期完成全部论文工作')
    if idx >= 0:
        insert_after(doc, idx,
            "从技术风险角度来看，系统的核心算法和架构已经过初步验证，不存在尚未攻克的关键技术难题。从实验条件来看，"
            "所需的大模型API服务和计算资源均已就位，具备开展实验的必要条件。从时间安排来看，当前距论文提交截止日期还有约11周时间，"
            "时间相对充裕，且已制定了详细的周计划，各项任务的时间估算合理。此外，导师在智能体和自然语言处理领域具有丰富的研究经验，"
            "能够在研究过程中提供及时的指导和建议。综上所述，按期完成全部论文工作具有较高的可行性。"
        )
        expansions.append("5 completion")
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ========== Final count ==========
    doc = Document(OUT_PATH)
    all_text = ''.join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    chinese_total = count_chinese(all_text)

    # Section 2 count
    sec2_text = ""
    started = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if '目前已经完成的研究工作' in txt:
            started = True
            continue
        if started and '后期拟完成的研究工作' in txt:
            break
        if started and txt:
            sec2_text += txt
    sec2_chinese = count_chinese(sec2_text)

    # Per-section counts
    sections = {}
    current_sec = 'header'
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        if '课题研究的背景和意义' in txt:
            current_sec = '1.1'
        elif '课题研究内容和进展' in txt:
            current_sec = '1.2'
        elif '目前已经完成的研究工作' in txt:
            current_sec = '2'
        elif '后期拟完成的研究工作' in txt:
            current_sec = '3'
        elif '存在的困难与问题' in txt:
            current_sec = '4'
        elif '如期完成全部论文工作的可能性' in txt:
            current_sec = '5'
        elif '参考文献' in txt and len(txt) < 10:
            current_sec = 'ref'
        if current_sec not in sections:
            sections[current_sec] = ''
        sections[current_sec] += txt

    print(f"\n{'='*40}")
    print(f"Expansions performed: {len(expansions)}")
    for e in expansions:
        print(f"  + {e}")
    print(f"\nPer-section Chinese chars:")
    for k in ['header', '1.1', '1.2', '2', '3', '4', '5', 'ref']:
        v = sections.get(k, '')
        c = count_chinese(v)
        print(f"  {k}: {c}")
    print(f"\nSection 2 Chinese: {sec2_chinese}")
    print(f"Total Chinese: {chinese_total}")
    print(f"Added: {chinese_total - initial_total}")
    print(f"{'='*40}")

    # Check for code identifiers
    bad_patterns = ['GeneratorFeedback', 'ValidatorFeedback', 'EvaluatorFeedback', 'cold_start', 'gen_only', 'val_only']
    found_any = False
    for p in bad_patterns:
        if p.lower() in all_text.lower():
            print(f'WARNING: Found "{p}"')
            found_any = True
    if not found_any:
        print("Code identifier check: PASSED")


if __name__ == '__main__':
    main()
