from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


REPORT_JSON = Path(
    "/Users/zhaoziqing/Desktop/benchforge/runs/effectiveness_task_20260618_153413/effectiveness/task_effectiveness_report.json"
)
OUT_DOC = Path("/Users/zhaoziqing/Desktop/benchforge/report/多轮实验结果与分析_独立版.docx")


def set_cn_font(run, *, size=12, bold=False):
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)


def add_para(doc: Document, text: str, *, size=12, bold=False, indent=True, space_after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    run = p.add_run(text)
    set_cn_font(run, size=size, bold=bold)
    return p


def add_heading(doc: Document, text: str, level: int):
    size_map = {1: 16, 2: 14, 3: 12.5}
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_cn_font(run, size=size_map.get(level, 12), bold=True)
    return p


def set_table_cell(cell, text: str, *, bold=False, size=10.5):
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            set_cn_font(r, size=size, bold=bold)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_table_cell(table.rows[0].cells[idx], header, bold=True)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            set_table_cell(table.rows[r_idx].cells[c_idx], value)
    return table


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def compute_verify_stats(task_dir: Path):
    result = {}
    for rid in ["round_001", "round_002"]:
        report = load_json(task_dir / rid / "validation" / "validation_report.json")
        citation_rows = [json.loads(line) for line in (task_dir / rid / "validation" / "citation_validation.jsonl").open() if line.strip()]
        llm_rows = [json.loads(line) for line in (task_dir / rid / "validation" / "llm_validation.jsonl").open() if line.strip()]
        result[rid] = {
            "report": report,
            "citation_mean": sum(r["citation_score"] for r in citation_rows) / len(citation_rows),
            "answer_citation_mean": sum(r["answer_citation_score"] for r in citation_rows) / len(citation_rows),
            "chunk_citation_mean": sum(r["chunk_citation_score"] for r in citation_rows) / len(citation_rows),
            "llm_mean": sum(r["overall_score"] for r in llm_rows) / len(llm_rows),
            "difficulty_consistency_mean": sum(r["dimensions"]["difficulty_consistency"] for r in llm_rows) / len(llm_rows),
            "latency_mean_ms": sum(r["latency_ms"] for r in llm_rows) / len(llm_rows),
            "input_token_mean": sum(r["input_tokens"] for r in llm_rows) / len(llm_rows),
            "output_token_mean": sum(r["output_tokens"] for r in llm_rows) / len(llm_rows),
        }
    return result


def main():
    data = load_json(REPORT_JSON)
    task_dir = Path(data["task_dir"])
    verify = compute_verify_stats(task_dir)

    b1 = data["b1"]
    bfull = data["bfull"]

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("多轮实验结果与分析")
    set_cn_font(run, size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于 effectiveness_task_20260618_153413 的独立实验说明")
    set_cn_font(run, size=12, bold=False)

    add_heading(doc, "1 实验目的与设置", 1)
    add_para(
        doc,
        "本实验用于验证当前自动评价总框架在真实多轮场景下的有效性，重点观察三个问题："
        "第一，多轮反馈是否能够在后续轮次继续提供高质量题目增量；第二，验证智能体是否能够稳定地完成事实性和质量性筛选；"
        "第三，模型评估智能体在新增题目加入后，是否能够给出更稳定、更有判别力的模型评分结果。实验在同一 task_id 下串行执行两轮实测，"
        "分别对应 round_001 与 round_002；每一轮为独立 run_id，从而避免不同轮次之间的文件覆盖。最终模型评测使用 deepseek-v4、"
        "deepseek-v4-pro 与 minimax 三个候选模型。",
    )

    add_heading(doc, "2 题目生成与总体结果", 1)
    add_para(
        doc,
        "两轮实测结果表明，多轮机制已经能够稳定提供题目增量。第1轮最终入选11道题，其中问答题8道、选择题3道；"
        "第2轮进一步入选13道题，其中问答题10道、选择题3道。两轮累计获得24道可用于后续模型评测的题目。"
        "从主题覆盖上看，第1轮更偏向 Artificial Intelligence，第2轮显著补充了 Quantum Computing 方向题目，说明多轮生成并不是重复第一轮内容，"
        "而是在后续轮次继续扩展相对薄弱的主题覆盖。",
    )
    add_table(
        doc,
        ["轮次", "入选题数", "QA题数", "选择题数", "主要主题倾向", "备注"],
        [
            ["第1轮（实测）", "11", "8", "3", "Artificial Intelligence 为主", "存在1道引文淘汰、1道LLM质量淘汰"],
            ["第2轮（实测）", "13", "10", "3", "Quantum Computing 为主", "13道题全部通过验证"],
            ["第3轮（趋势预测，非实测）", "12-14", "9-11", "3-4", "继续补齐薄弱主题与难度层", "若引入区分度反馈，预计更强调高判别题目"],
        ],
    )
    add_para(
        doc,
        "第三轮数据并非实测结果，而是基于前两轮趋势给出的预测性判断，仅用于说明后续实验预期方向，不能与前两轮实测结果等量齐观。"
        "如果第三轮继续沿用当前策略，题目总数预计仍会增长，但增长幅度可能趋于平稳；若在反馈中显式加入区分度信号，则第三轮更可能优先生成"
        "对模型能力边界更敏感的题目。",
    )

    add_heading(doc, "3 验证智能体实验分析", 1)
    add_para(
        doc,
        "验证智能体的表现总体较为稳定，并且第2轮相较第1轮表现出明显改善。第1轮共有13道候选题，12道通过引文验证，11道通过大模型质量验证，"
        "最终入选11道，说明验证智能体在首轮能够有效过滤掉事实支撑不足或整体质量不达标的题目。第2轮同样接收13道候选题，但13道题全部通过引文验证、"
        "全部通过大模型质量验证并被最终选中，表明在后续轮次中，生成策略与验证标准之间的匹配程度更高。",
    )
    add_table(
        doc,
        ["轮次", "候选题数", "引文通过", "LLM通过", "最终入选", "引文均值", "LLM均值"],
        [
            [
                "第1轮",
                str(verify["round_001"]["report"]["total_candidates"]),
                str(verify["round_001"]["report"]["citation_passed"]),
                str(verify["round_001"]["report"]["llm_passed"]),
                str(verify["round_001"]["report"]["final_selected"]),
                f"{verify['round_001']['citation_mean']:.4f}",
                f"{verify['round_001']['llm_mean']:.4f}",
            ],
            [
                "第2轮",
                str(verify["round_002"]["report"]["total_candidates"]),
                str(verify["round_002"]["report"]["citation_passed"]),
                str(verify["round_002"]["report"]["llm_passed"]),
                str(verify["round_002"]["report"]["final_selected"]),
                f"{verify['round_002']['citation_mean']:.4f}",
                f"{verify['round_002']['llm_mean']:.4f}",
            ],
        ],
    )
    add_para(
        doc,
        f"从引文验证细项来看，第1轮平均引文得分为 {verify['round_001']['citation_mean']:.4f}，第2轮提升到 "
        f"{verify['round_002']['citation_mean']:.4f}；其中答案引文支撑均值从 {verify['round_001']['answer_citation_mean']:.4f} "
        f"提升到 {verify['round_002']['answer_citation_mean']:.4f}，文本块匹配均值则从 {verify['round_001']['chunk_citation_mean']:.4f} "
        f"提升到 {verify['round_002']['chunk_citation_mean']:.4f}。这说明第二轮不仅“通过率更高”，而且引用质量本身也更强，尤其是答案与证据的对齐程度进一步改善。",
    )
    add_para(
        doc,
        f"大模型质量验证也呈现相同趋势。第1轮 LLM 综合质量分均值为 {verify['round_001']['llm_mean']:.4f}，"
        f"第2轮提升到 {verify['round_002']['llm_mean']:.4f}；难度一致性维度均值从 "
        f"{verify['round_001']['difficulty_consistency_mean']:.4f} 提升到 {verify['round_002']['difficulty_consistency_mean']:.4f}。"
        "这表明后续轮次生成的题目在“标注难度是否与题目真实复杂度一致”这一点上更加稳定。"
        "值得注意的是，第2轮日志中出现过一次“Cannot parse JSON from LLM response”的告警，但随后重试成功，最终13/13全部通过，"
        "说明当前验证智能体已经具备一定的容错和自动恢复能力。",
    )
    add_para(
        doc,
        f"从资源消耗看，第1轮验证阶段共进行12次 LLM 调用，输入/输出 token 分别为 "
        f"{verify['round_001']['report']['llm_usage']['input_tokens']} 和 {verify['round_001']['report']['llm_usage']['output_tokens']}；"
        f"第2轮为13次调用，输入/输出 token 分别为 {verify['round_002']['report']['llm_usage']['input_tokens']} 和 "
        f"{verify['round_002']['report']['llm_usage']['output_tokens']}。虽然第2轮调用量略有增加，但平均输出 token 从 "
        f"{verify['round_001']['output_token_mean']:.1f} 下降到 {verify['round_002']['output_token_mean']:.1f}，"
        "说明模型在后续轮次的验证回答更集中、更稳定，质量提升并未伴随明显的额外裁判冗余。",
    )

    add_heading(doc, "4 模型评估智能体实验分析", 1)
    add_para(
        doc,
        "模型评估部分分别在两个题集上进行：B1 表示仅使用第1轮入选题（11题），Bfull 表示使用两轮累计入选题（24题）。"
        "这两个题集的比较能够帮助判断后续轮次新增题目是否仅带来数量扩张，还是同时影响了模型区分度与数据集质量。",
    )
    add_table(
        doc,
        ["题集", "题目数", "引用得分", "多样性得分", "Top-Bottom Gap", "方差", "区分题比例"],
        [
            [
                "B1（第1轮）",
                str(b1["question_count"]),
                f"{b1['evaluation_result']['dataset_metrics']['citation_score']['mean']:.4f}",
                f"{b1['evaluation_result']['dataset_metrics']['diversity_score']['score']:.4f}",
                f"{b1['benchmark_metrics']['top_bottom_gap']:.4f}",
                f"{b1['benchmark_metrics']['score_variance']:.6f}",
                f"{b1['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:.4f}",
            ],
            [
                "Bfull（两轮累计）",
                str(bfull["question_count"]),
                f"{bfull['evaluation_result']['dataset_metrics']['citation_score']['mean']:.4f}",
                f"{bfull['evaluation_result']['dataset_metrics']['diversity_score']['score']:.4f}",
                f"{bfull['benchmark_metrics']['top_bottom_gap']:.4f}",
                f"{bfull['benchmark_metrics']['score_variance']:.6f}",
                f"{bfull['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:.4f}",
            ],
        ],
    )
    add_para(
        doc,
        "首先，从数据集自身质量看，后续轮次新增题目带来了正向增益。Bfull 的引用得分由 0.7894 提升到 0.8066，"
        "多样性得分由 0.8232 提升到 0.8912，说明后续轮次在事实支撑和语义分散性上均对题集形成了有效补充。"
        "这意味着多轮反馈已经不仅仅是在“堆数量”，而是在改善最终评测数据的覆盖范围和资料质量。",
    )
    add_para(
        doc,
        f"其次，从总体模型排序看，B1 上的综合得分排序为 minimax（{b1['benchmark_metrics']['model_scores']['minimax']:.4f}）"
        f"> deepseek-v4（{b1['benchmark_metrics']['model_scores']['deepseek-v4']:.4f}）"
        f"> deepseek-v4-pro（{b1['benchmark_metrics']['model_scores']['deepseek-v4-pro']:.4f}）；"
        f"Bfull 上仍为 minimax（{bfull['benchmark_metrics']['model_scores']['minimax']:.4f}）"
        f"> deepseek-v4（{bfull['benchmark_metrics']['model_scores']['deepseek-v4']:.4f}）"
        f"> deepseek-v4-pro（{bfull['benchmark_metrics']['model_scores']['deepseek-v4-pro']:.4f}）。"
        "然而，如果仅看问答题（QA）子集，排序则完全相反：deepseek-v4 始终最高，deepseek-v4-pro 次之，minimax 最低。"
        "这说明整体排序实际上受到选择题子集聚合方式的强烈影响，不能直接作为三模型真实综合能力的最终结论。",
    )
    add_para(
        doc,
        "从 QA 自动指标细项看，deepseek-v4 在两种题集上都保持了较强的稳定性：其 semantic accuracy 从 0.375 提升到 0.4444，"
        "F1、precision、ROUGE-L 和 BLEU 均保持领先或接近领先；deepseek-v4-pro 在语义准确率和 BERTScore 上相对稳定，但整体生成质量略低于 deepseek-v4；"
        "minimax 在 recall 上异常偏高（B1 为 0.8296，Bfull 为 0.7624），但 precision 明显偏低（分别仅为 0.1383 和 0.1585），"
        "说明其回答存在较强的“覆盖式作答”倾向，即输出内容较多、命中部分参考答案成分，但精确对齐能力不足。",
    )
    add_para(
        doc,
        "LLM 裁判指标进一步强化了上述判断。在 B1 上，deepseek-v4 与 deepseek-v4-pro 在 answer relevance、faithfulness、factual accuracy、"
        "logical coherence 和 completeness 五个维度上基本保持满分；minimax 则降至 4.0 左右。到了 Bfull，deepseek-v4 和 deepseek-v4-pro 仍保持 5.0 左右的稳定表现，"
        "而 minimax 在五个维度上进一步下降到 3.0-3.625 区间，说明新增题目对 minimax 的鲁棒性提出了更高要求。这一现象与 QA 自动指标中 minimax 的 precision 偏低、"
        "semantic accuracy 偏低相吻合，表明其在开放式问答上的真实表现弱于另外两种模型。",
    )
    add_para(
        doc,
        "选择题（multiple-choice）子集暴露了当前模型评估聚合的一项重要问题。自动指标显示，在 B1 与 Bfull 上，deepseek-v4 和 deepseek-v4-pro 的选择题 accuracy 均为 1.0，"
        "invalid rate 为 0 或 0.1667；而 minimax 的 invalid rate 始终为 1.0，且 accuracy 无法计算。按常理，这意味着 minimax 在选择题输出格式上表现最差。"
        "然而在综合报告中，multiple-choice 模式下 minimax 的模型得分却显示为 1.0，高于另外两个模型的 0.5 或 0.5455。"
        "这表明当前模型评估智能体在多选题指标聚合时，尚未对“正向指标”和“反向指标”进行统一方向化处理，导致 invalid rate 这样的反向指标可能被错误地当作正向贡献。"
        "因此，当前综合得分中 minimax 的领先不应被直接解释为其真实能力领先，而更可能是评测聚合实现中的方向性缺陷所致。",
    )
    add_para(
        doc,
        "从区分度信号看，Bfull 的区分题比例由 0.4545 小幅上升到 0.4583，说明新增题目并未削弱题集的基本判别能力；但 overall model gap、best-vs-second gap、"
        "top-bottom gap 和 per-question variance 均略有下降，说明两轮累计题集更多提升了覆盖面与多样性，却没有同步显著放大模型间总体分差。"
        "换言之，当前多轮策略已经证明其能“把题集做大、做稳、做多样”，但尚未充分证明其能“让题集更会区分模型”。"
    )

    add_heading(doc, "5 第三轮趋势预测（非实测）", 1)
    add_para(
        doc,
        "结合前两轮趋势，如果继续采用当前机制而不显式引入区分度反馈，第三轮预计仍会产生一定数量的新增入选题，引用得分和多样性可能继续缓慢上升，"
        "但模型间总体分差未必同步扩大。更合理的预期是：第三轮应将重点从“继续扩题”转向“筛选高判别题”，即把模型分差、题目方差、裁判分歧度、"
        "以及多选题格式错误率等信号接入反馈环路，优先生成能够稳定拉开 deepseek-v4、deepseek-v4-pro 与 minimax 差异的题目。"
        "因此，本文建议在后续实验中将第三轮设计为“区分度导向轮”：减少单纯覆盖式扩展，转而强化对高信息量题目的定向搜索。",
    )

    add_heading(doc, "6 结论与可写入中期报告的要点", 1)
    add_para(
        doc,
        "综合本次独立实验可以得出四点结论。第一，多轮反馈机制已经能够在后续轮次稳定提供题目增量，且新增题目在引用质量和多样性上优于首轮题集。"
        "第二，验证智能体在第2轮表现出更高的稳定性和更强的质量一致性，证明当前验证链路具备较好的容错和收敛特征。第三，模型评估智能体在 QA 子集上已经能够稳定给出较为一致的模型排序，"
        "但在选择题聚合上存在反向指标未做方向统一的问题，这一实现缺陷会影响综合模型得分解释。第四，从真实研究价值看，当前框架已经证明了“数据生成与验证闭环”的有效性，"
        "下一阶段工作的重点应当从题目数量扩展转向区分度增强和评测聚合修正。",
    )

    doc.save(OUT_DOC)


if __name__ == "__main__":
    main()
