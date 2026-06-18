"""Final expansion to push past 10000 Chinese characters."""
from docx import Document
from lxml import etree
from docx.oxml.ns import qn

OUT_PATH = '/Users/zhaoziqing/Desktop/benchforge/report/赵子晴_中期报告_完善版.docx'

FONT_BODY = 'Times New Roman'
FONT_EA = '宋体'
SZ_BODY = 12
ALIGN_JUSTIFY = 'both'


def make_para(doc, text):
    new_p = etree.Element(qn('w:p'))
    pPr = etree.SubElement(new_p, qn('w:pPr'))
    jc = etree.SubElement(pPr, qn('w:jc'))
    jc.set(qn('w:val'), ALIGN_JUSTIFY)
    r = etree.SubElement(new_p, qn('w:r'))
    rPr = etree.SubElement(r, qn('w:rPr'))
    rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    rFonts.set(qn('w:ascii'), FONT_BODY)
    rFonts.set(qn('w:hAnsi'), FONT_BODY)
    rFonts.set(qn('w:eastAsia'), FONT_EA)
    sz = etree.SubElement(rPr, qn('w:sz'))
    sz.set(qn('w:val'), str(int(SZ_BODY * 2)))
    szCs = etree.SubElement(rPr, qn('w:szCs'))
    szCs.set(qn('w:val'), str(int(SZ_BODY * 2)))
    t = etree.SubElement(r, qn('w:t'))
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    return new_p


def insert_after(doc, after_idx, text):
    ref_element = doc.paragraphs[after_idx]._element
    new_p = make_para(doc, text)
    ref_element.addnext(new_p)


def find_text(doc, keyword, start=0):
    for i in range(start, len(doc.paragraphs)):
        if keyword in doc.paragraphs[i].text.strip():
            return i
    return -1


def count_chinese(text):
    return sum(1 for c in text if '一' <= c <= '鿿')


doc = Document(OUT_PATH)
initial_total = count_chinese(''.join(p.text.strip() for p in doc.paragraphs if p.text.strip()))
print('Initial total Chinese:', initial_total)

expansions = []

# === Section 4: expand compute resource discussion ===
idx = find_text(doc, '计算资源消耗问题')
if idx >= 0:
    text = (
        '针对计算资源消耗问题，本研究在实践过程中探索了多种优化策略。'
        '在模型调用层面，通过引入批量请求和连接复用机制，降低了API调用的单位开销。'
        '在任务调度层面，采用优先级队列对生成任务进行排序，优先执行对整体进度影响最大的关键任务，'
        '避免在低价值任务上过度消耗资源。'
        '在数据复用层面，已验证通过的题目及其验证结果会被缓存，'
        '同一题目在后续轮次中无需重新验证，减少了重复计算。'
        '这些优化策略在实际运行中取得了较好的效果，'
        '但如何在评价质量和计算成本之间找到最佳平衡点，'
        '仍是一个需要持续探索的问题。'
    )
    insert_after(doc, idx, text)
    expansions.append('4 compute optimization')
    doc.save(OUT_PATH)
    doc = Document(OUT_PATH)

# === Section 4: expand difficulty about benchmark generalization ===
idx = find_text(doc, '评价结果的稳定性和可复现性问题')
if idx >= 0:
    text = (
        '此外，评价数据集的泛化能力也是一个值得关注的问题。'
        '当前系统生成的评价数据集主要基于百科类知识源，'
        '虽然覆盖了多个主题领域，但这些知识源的文本风格和内容组织方式相对单一，'
        '可能导致生成的题目在某些方面存在系统性偏差。'
        '例如，对于需要考察模型实际应用能力的场景，'
        '基于百科知识生成的题目可能无法充分反映真实使用环境中的复杂性。'
        '未来工作将探索引入更多样化的知识来源，'
        '包括技术文档、学术论文和实际应用案例等，'
        '以提升评价数据集的覆盖面和生态效度。'
        '同时，也需要研究不同知识源对评价结果的影响，'
        '建立知识源质量评估和筛选机制，确保输入资料的质量和适宜性。'
    )
    insert_after(doc, idx, text)
    expansions.append('4 generalization')
    doc.save(OUT_PATH)
    doc = Document(OUT_PATH)

# === Final count ===
doc = Document(OUT_PATH)
all_text = ''.join(p.text.strip() for p in doc.paragraphs if p.text.strip())
chinese_total = count_chinese(all_text)

# Per-section counts
sections = {}
current_sec = 'header'
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    if '课题研究的背景和意义' in txt:
        current_sec = '1.1'
    elif '课题研究内容和进展' in txt:
        current_sec = '1.2'
    elif '目前已经完成的研究工作' in txt:
        current_sec = '2'
    elif '后期拟完成的研究工作' in txt:
        current_sec = '3'
    elif '存在的困难与问题' in txt:
        current_sec = '4'
    elif '如期完成全部论文工作的可能性' in txt:
        current_sec = '5'
    elif txt == '参考文献':
        current_sec = 'ref'
    if current_sec not in sections:
        sections[current_sec] = ''
    sections[current_sec] += txt

print('\n' + '=' * 40)
print('Expansions performed:', len(expansions))
for e in expansions:
    print('  +', e)
print('\nPer-section Chinese chars:')
for k in ['header', '1.1', '1.2', '2', '3', '4', '5', 'ref']:
    v = sections.get(k, '')
    c = count_chinese(v)
    print('  {}: {}'.format(k, c))
print('\nTotal Chinese:', chinese_total)
print('Added:', chinese_total - initial_total)
print('=' * 40)

# Check for code identifiers
bad_patterns = ['GeneratorFeedback', 'ValidatorFeedback', 'EvaluatorFeedback',
                'cold_start', 'gen_only', 'val_only']
found_any = False
for p in bad_patterns:
    if p.lower() in all_text.lower():
        print('WARNING: Found', p)
        found_any = True
if not found_any:
    print('Code identifier check: PASSED')

if chinese_total >= 10000:
    print('=== TARGET REACHED: 10000+ Chinese characters ===')
else:
    print('Still need', 10000 - chinese_total, 'more characters')
