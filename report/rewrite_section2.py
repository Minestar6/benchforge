"""Rewrite section 2 of the mid-term report with user's required structure.

Structure:
  2.1 评价数据集构造: (1)证据池构造 (2)题目生成介绍 (3)多轮反馈机制
  2.2 自动评价智能体架构: (1)规划智能体 (2)题目验证智能体 (3)模型评估智能体 (4)总体框架介绍
  2.3 评估报告生成

Content must use natural technical language - NO code-level identifiers like GeneratorFeedback, cold_start, etc.
"""

from docx import Document
from lxml import etree
from docx.oxml.ns import qn

OUT_PATH = '/Users/zhaoziqing/Desktop/benchforge/report/赵子晴_中期报告_完善版.docx'

FONT_BODY = 'Times New Roman'
FONT_EA = '宋体'
FONT_HEADING_EA = '黑体'
SZ_BODY = 12
SZ_SUBSECTION = 14
ALIGN_JUSTIFY = 'both'


def make_para(doc, text, bold=False, size=SZ_BODY, font_name=FONT_BODY,
              ea_font_name=FONT_EA, alignment=ALIGN_JUSTIFY):
    new_p = etree.Element(qn('w:p'))
    pPr = etree.SubElement(new_p, qn('w:pPr'))
    if alignment is not None:
        jc = etree.SubElement(pPr, qn('w:jc'))
        jc.set(qn('w:val'), alignment)
    r = etree.SubElement(new_p, qn('w:r'))
    rPr = etree.SubElement(r, qn('w:rPr'))
    rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    if font_name:
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
    if ea_font_name:
        rFonts.set(qn('w:eastAsia'), ea_font_name)
    if size:
        sz = etree.SubElement(rPr, qn('w:sz'))
        sz.set(qn('w:val'), str(int(size * 2)))
        szCs = etree.SubElement(rPr, qn('w:szCs'))
        szCs.set(qn('w:val'), str(int(size * 2)))
    if bold:
        etree.SubElement(rPr, qn('w:b'))
        etree.SubElement(rPr, qn('w:bCs'))
    t = etree.SubElement(r, qn('w:t'))
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    return new_p


def insert_many(doc, after_idx, items):
    """Insert multiple paragraphs after the given index in order.

    items: list of dicts with keys: text, bold, size
    """
    ref_element = doc.paragraphs[after_idx]._element
    for item in reversed(items):
        new_p = make_para(
            doc,
            text=item['text'],
            bold=item.get('bold', False),
            size=item.get('size', SZ_BODY),
            font_name=item.get('font_name', FONT_BODY),
            ea_font_name=item.get('ea_font_name', FONT_EA),
            alignment=item.get('alignment', ALIGN_JUSTIFY),
        )
        ref_element.addnext(new_p)


# ========== CONTENT ==========

sec2_1_heading = "2.1 评价数据集构造"

sec2_1_sub1_heading = "（1）证据池构造"
sec2_1_sub1_body = (
    "证据池构造是题目生成的基础环节。系统首先根据用户指定的研究主题，通过自动化检索手段从百科类知识源中获取相关原始文档。"
    "获取到原始文档后，系统对文档内容进行结构化处理，采用可配置的文本分段策略将长文档切分为语义完整的文本块。系统支持两种分段策略："
    "基于滑动窗口的固定长度分块，将文档按指定长度均匀切分，窗口大小和重叠量均可配置；基于语义边界的自适应分块，依据文档的自然段落结构和"
    "语义完整性进行分段，保留原文的章节层次关系。每个文本块均携带来源元数据，包括所属文档标题、章节信息和位置索引，确保生成的题目具有"
    "完整的可追溯性。为进一步提升证据的利用效率，系统还构建了单文本块和多文本块两种粒度的证据单元：单文本块单元直接基于单个分段构造题目，"
    "适用于考察细粒度知识点；多文本块单元则组合多个相关文本块，适用于考察跨段落的综合理解能力。同时，系统调用大模型为每个文档生成简短的摘要，"
    "便于后续的主题筛选和证据优先级排序。整个证据池构造过程完全自动化，用户只需指定研究主题即可自动完成相关资料的检索和预处理。"
)

sec2_1_sub2_heading = "（2）题目生成介绍"
sec2_1_sub2_body = (
    "题目生成过程以证据池为基础，按照用户指定的题目模式和难度分布，系统性地调用大模型生成候选题目。系统当前支持两种题目模式：问答题要求模型"
    "根据证据片段生成问题及其标准答案，并标注答案在证据中的引用位置；多项选择题则在生成正确答案的同时构造具有迷惑性的干扰项，确保所有选项"
    "均与证据内容相关。具体生成流程为：首先根据当前轮的生成策略和难度分布，从证据池中采样合适的证据单元。采样过程采用预分配机制，在正式调用"
    "大模型之前预先为每个主题和难度层级分配好需要使用的证据单元，避免不同主题之间的资源竞争。随后，系统对每个采样的证据单元调用一次大模型，"
    "基于其中的文本内容按照预设的提示词模板生成一道或多道题目。每道生成的题目包含丰富的信息字段：题目文本、标准答案（多项选择题还包括选项列表"
    "和正确选项）、预估难度等级、所属主题、证据引用信息（引用文本块的具体内容和位置）、题目类型、所需能力标签以及模型的推理过程。生成完成后，"
    "系统通过轻量级过滤器对题目进行初步筛选，剔除格式异常、内容空洞或明显错误的题目，并通过基于哈希值的精确去重机制移除完全重复的题目。"
    "初步筛选后的合格题目进入候选池，等待后续的多阶段验证。"
)

sec2_1_sub3_heading = "（3）多轮反馈机制"
sec2_1_sub3_body = (
    "题目生成并非一次完成，而是通过多轮迭代逐步优化。系统设计了阶段化的生成策略演化路径：初始阶段以广度覆盖为主，尽量覆盖全部主题和难度层级，"
    "快速积累基础题目数量；在积累一定数量的题目后，逐步转向针对性生成，重点补充困难题目和表现薄弱的主题方向；当现有证据已充分使用时，触发"
    "证据扩展流程，为目标主题检索更多相关资料以扩展生成空间；最后阶段进行终末修复，针对仍然不足的难度层级进行定向补充。每一轮生成结束后，"
    "系统会综合分析本轮的各难度层级产量、连续空轮次数、生成失败率等信息，作为下一轮策略调整的依据。在证据采样方面，系统会跟踪每个文本块在"
    "各轮中的使用频率，优先选择尚未充分利用的证据片段，保证证据资源的均匀使用和题目多样性。系统还通过产量估计机制预测当前证据资源的剩余生成"
    "潜力，避免在低效证据上继续浪费计算资源。整个多轮生成过程受多种停止条件约束：当候选池达到目标数量上限、达到最大生成轮数上限、连续多轮"
    "无新题目生成或生成失败率超过阈值时，系统自动终止生成。通过这些机制，系统能够在无需人工干预的情况下，自主完成大规模、高质量评价数据集的"
    "构建，并具有良好的鲁棒性和自适应能力。"
)

sec2_2_heading = "2.2 自动评价智能体架构"

sec2_2_sub4_heading = "（4）总体框架介绍"
sec2_2_sub4_body = (
    "本系统采用强中心化的多智能体架构，由一个核心的规划智能体统一调度三个专门化的子智能体：题目生成智能体负责构建评价数据集、题目验证智能体"
    "负责质量把关和筛选、模型评估智能体负责在数据集上对候选模型进行系统评估。这种中心化架构的核心优势在于决策的全局一致性——所有智能体的行为"
    "均由规划智能体根据用户设定的整体目标和当前系统状态统一协调，避免了分布式决策可能导致的资源冲突和策略不一致问题。系统的工作流程是一个"
    "动态的多轮反馈闭环：每一轮迭代中，规划智能体首先诊断上一轮各智能体返回的执行结果，综合分析当前进度和存在的问题，制定本轮的具体执行计划；"
    "然后依次调用各子智能体执行具体任务；执行完成后收集各环节的反馈信息，更新系统状态，进入下一轮的诊断和规划。这种动态反馈机制使得系统能够"
    "根据实际执行情况自适应地调整策略，逐步逼近用户的评价目标。整个架构通过编排器组件实现各组件的解耦与协作，每个子智能体专注于自身的核心"
    "职责，通过标准化的接口与规划智能体交互，具有良好的模块化程度和可扩展性。"
)

sec2_2_sub1_heading = "（1）规划智能体"
sec2_2_sub1_body = (
    "规划智能体是整个系统的核心决策组件，负责将用户的高层评价目标转化为可逐步执行的具体计划，并在多轮执行过程中持续监控进度和调整策略。"
    "其工作流程为诊断、规划、执行、反馈的闭环迭代过程。在每轮迭代开始时，规划智能体首先执行诊断阶段，综合分析上一轮生成的产量数据、验证通过率"
    "和质量评分、以及模型评估的得分分布，判断当前所处的阶段状态，并识别存在的问题（如题目区分度不足、质量偏低、分布失衡等）。基于诊断结果，"
    "规划智能体从三个维度制定本轮计划：第一，题目计划，确定各模式和难度层级的目标产量，以及本轮需要重点关注的主题列表，系统支持基于规则的"
    "固定演化策略和基于大模型的智能演化策略两种互补模式；第二，控制计划，配置验证环节的严格程度（包括引文验证的最低分数阈值、选择模式等）"
    "和评估环节的详细程度；第三，主题自适应选择，根据历史数据中各主题的模型评分差异、候选数量、质量指标等表现，动态筛选出当前最需要关注的"
    "主题方向。完成计划制定后，规划智能体将上述计划组装为可执行的轮次规范，分发至各子智能体执行，并在执行完成后收集各环节的反馈信息，更新"
    "系统状态，进入下一轮迭代。如此循环往复，直至达到用户的评价目标或满足终止条件。"
)

sec2_2_sub2_heading = "（2）题目验证智能体"
sec2_2_sub2_body = (
    "题目验证智能体负责对生成智能体产出的候选题目进行多阶段质量把关，确保最终纳入评价集的题目具有高质量、多样性和可评价性。验证流程分为"
    "三个阶段。第一阶段为引文验证：对每道题目引用的证据来源进行逐条核查，判断引文内容是否真实支持题目中的陈述和标准答案，并记录每条引文的"
    "支持强度和具体的支持原文，为后续的加权筛选提供依据。此阶段还同时执行精确去重，基于题目文本的哈希值移除完全重复的题目。第二阶段为大模型"
    "质量评估：调用大模型对通过引文验证的候选题目进行多维度的质量评分，评估维度包括题目表述的清晰度、标准答案的正确性、引文与题目的一致性、"
    "难度设置的合理性等。在进入此阶段前，系统还执行语义近似去重——通过计算题目文本嵌入向量之间的余弦相似度，移除语义高度重复的题目，确保"
    "评价集的多样性和覆盖广度。第三阶段为加权选择：综合引文验证分数和大模型质量评分，计算每道题目的综合权重，按照题目模式和难度层级进行"
    "等比例采样，确保最终入选的题目在不同模式、不同难度层级和不同主题之间保持合理的分布比例。选定的题目被标记为选中状态，并记录完整的验证"
    "信息备查。规划智能体可以调控该组件的多项关键参数，包括引文验证的各项分数阈值、选择模式的严格程度、语义相似度阈值等，并根据验证反馈"
    "（各阶段通过率、质量评分分布、去重比例、多样性指标等）动态调整这些参数，实现验证策略的自适应优化。"
)

sec2_2_sub3_heading = "（3）模型评估智能体"
sec2_2_sub3_body = (
    "模型评估智能体负责在构建好的评价数据集上对候选大模型进行系统性的能力评估，自动生成多维度、多粒度的评估结果。评估流程分为四个步骤。"
    "第一步，数据集质量评估：对评价数据集本身进行质量分析，计算引文覆盖率（衡量题目答案是否有可靠的引文支撑）和多样性分数（通过嵌入向量的"
    "离散度和聚类熵评估题目覆盖的广度），从源头保证评估的有效性。第二步，候选模型推理：将评价集中的题目逐一输入待评估的候选大模型，收集各模型"
    "生成的回答。系统通过统一的模型客户端接口抽象不同模型的调用差异，支持多种模型的并行推理，显著提升评估效率。第三步，自动指标评分：对模型"
    "回答进行客观的自动化指标评估，支持的指标包括精确匹配、关键词匹配和语义相似度等，不同题目模式可配置不同的指标组合和权重。第四步，大模型"
    "裁判评估：调用大模型作为裁判，对模型回答进行更深层次的质性分析，从回答的正确性、完整性、逻辑一致性和推理质量等多个维度进行综合打分。"
    "四个步骤的评估结果被统一聚合，生成结构化的综合评估报告，包含按题目模式、主题领域、难度层级和候选模型的多维度指标分解，以及模型间的"
    "区分度分析（包括模型得分差距、各题目的得分方差、能有效区分模型能力的题目比例等），为研究人员提供全方位、多层次的模型能力分析视图。"
    "规划智能体可以调控该组件的评估配置，并根据评估反馈（模型得分分布、区分度指标等）调整下一轮的生成策略和验证策略，形成全链路的自适应优化。"
)

sec2_3_heading = "2.3 评估报告生成"
sec2_3_body = (
    "评估报告生成是本系统的重要组成部分，其目标是将模型评估智能体产生的结构化评估数据转化为直观、全面的综合评价报告。在评估数据的组织与"
    "聚合方面，系统设计了一套层次化的数据结构，能够将原始评估指标从按题目、按模式、按主题、按难度、按模型的多个维度进行有效整合，形成"
    "结构清晰、层次分明的评估结果呈现。在评估结果的分析与解读方面，系统自动计算一系列高级分析信号，包括候选模型之间的总体能力差距、领先模型"
    "的优势程度、各题目在不同模型上的得分离散度、以及评价数据集本身的区分度质量指标（如能有效区分模型能力的题目比例、过易题目比例、全部模型"
    "均失败的题目比例等），为评价数据集的质量诊断提供量化依据。在具体实现上，系统目前已实现了评估报告的标准化数据结构输出，能够生成包含多维"
    "指标分解和区分度分析的评估结果。在此基础上，后续将进一步开发面向人类研究者的可读报告生成能力，包括自动生成评估结果的文本摘要和结论分析、"
    "绘制雷达图和柱状图等可视化图表，以及基于评估数据自动提炼各模型的优势领域和不足之处，帮助研究人员快速理解各候选模型的能力差异，为模型"
    "选型和改进提供数据驱动的决策支持。"
)


def main():
    doc = Document(OUT_PATH)

    # Find section boundaries
    sec2_heading_idx = -1
    sec3_heading_idx = -1
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if '目前已经完成的研究工作' in txt:
            sec2_heading_idx = i
        if sec2_heading_idx >= 0 and '后期拟完成的研究工作' in txt:
            sec3_heading_idx = i
            break

    print(f"sec2 heading: {sec2_heading_idx}, sec3 heading: {sec3_heading_idx}")

    if sec2_heading_idx < 0 or sec3_heading_idx < 0:
        print("ERROR: Could not find section boundaries")
        return

    # Remove existing section 2 content (between sec2 heading and sec3 heading)
    elements_to_remove = []
    for i in range(sec2_heading_idx + 1, sec3_heading_idx):
        elements_to_remove.append(doc.paragraphs[i]._element)

    for elem in elements_to_remove:
        parent = elem.getparent()
        if parent is not None:
            parent.remove(elem)

    print(f"Removed {len(elements_to_remove)} paragraphs")

    # Save and reload after removal
    doc.save(OUT_PATH)
    doc = Document(OUT_PATH)

    # Re-find sec2 heading index after removal
    sec2_heading_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if '目前已经完成的研究工作' in p.text.strip():
            sec2_heading_idx = i
            break

    print(f"New sec2 heading index: {sec2_heading_idx}")

    # Insert new content in reverse order (each insertion goes after the heading)
    # Build all paragraphs in display order (top-to-bottom)
    all_items = [
        # 2.1 评价数据集构造
        {'text': sec2_1_heading, 'bold': True, 'size': SZ_SUBSECTION, 'ea_font_name': FONT_HEADING_EA},
        {'text': sec2_1_sub1_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_1_sub1_body, 'bold': False, 'size': SZ_BODY},
        {'text': sec2_1_sub2_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_1_sub2_body, 'bold': False, 'size': SZ_BODY},
        {'text': sec2_1_sub3_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_1_sub3_body, 'bold': False, 'size': SZ_BODY},

        # 2.2 自动评价智能体架构
        {'text': sec2_2_heading, 'bold': True, 'size': SZ_SUBSECTION, 'ea_font_name': FONT_HEADING_EA},
        {'text': sec2_2_sub1_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_2_sub1_body, 'bold': False, 'size': SZ_BODY},
        {'text': sec2_2_sub2_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_2_sub2_body, 'bold': False, 'size': SZ_BODY},
        {'text': sec2_2_sub3_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_2_sub3_body, 'bold': False, 'size': SZ_BODY},
        {'text': sec2_2_sub4_heading, 'bold': True, 'size': SZ_BODY},
        {'text': sec2_2_sub4_body, 'bold': False, 'size': SZ_BODY},

        # 2.3 评估报告生成
        {'text': sec2_3_heading, 'bold': True, 'size': SZ_SUBSECTION, 'ea_font_name': FONT_HEADING_EA},
        {'text': sec2_3_body, 'bold': False, 'size': SZ_BODY},
    ]

    insert_many(doc, sec2_heading_idx, all_items)
    doc.save(OUT_PATH)

    # === Verify ===
    doc = Document(OUT_PATH)
    sec2_text = ""
    started = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if '目前已经完成的研究工作' in txt:
            started = True
            continue
        if started and '后期拟完成的研究工作' in txt:
            break
        if started and txt:
            sec2_text += txt

    chinese = sum(1 for c in sec2_text if '一' <= c <= '鿿')
    total_chinese = sum(1 for p in doc.paragraphs for c in p.text.strip() if '一' <= c <= '鿿')
    print(f"\nSection 2: {chinese} Chinese characters, {len(sec2_text)} total chars")
    print(f"Total document: {total_chinese} Chinese characters, {len(doc.paragraphs)} paragraphs")

    # Print structure
    print("\nSection 2 structure:")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        if i >= sec2_heading_idx or ('目前已经完成的研究工作' in txt):
            if any(kw in txt for kw in ['2.1', '2.2', '2.3', '目前已经完成', '后期拟完成']):
                print(f"  [{i}] {txt[:80]}")
                if '后期拟完成' in txt:
                    break
            elif txt.startswith('（'):
                print(f"  [{i}]   {txt[:80]}")

    print(f"\nSaved to: {OUT_PATH}")


if __name__ == '__main__':
    main()
