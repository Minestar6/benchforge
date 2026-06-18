# -*- coding: utf-8 -*-
"""Expand the mid-term report to 10000+ Chinese characters."""

from docx import Document
from lxml import etree
from docx.oxml.ns import qn

OUT_PATH = '/Users/zhaoziqing/Desktop/benchforge/report/赵子晴_中期报告_完善版.docx'

FONT_BODY = 'Times New Roman'
FONT_EA = '宋体'
FONT_HEADING_EA = '黑体'
SZ_BODY = 12
SZ_SUBSECTION = 14
ALIGN_JUSTIFY = 'both'


def make_para(doc, text, bold=False, size=SZ_BODY, font_name=FONT_BODY,
              ea_font_name=FONT_EA, alignment=ALIGN_JUSTIFY):
    new_p = etree.Element(qn('w:p'))
    pPr = etree.SubElement(new_p, qn('w:pPr'))
    if alignment is not None:
        jc = etree.SubElement(pPr, qn('w:jc'))
        jc.set(qn('w:val'), alignment)
    r = etree.SubElement(new_p, qn('w:r'))
    rPr = etree.SubElement(r, qn('w:rPr'))
    rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    if font_name:
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
    if ea_font_name:
        rFonts.set(qn('w:eastAsia'), ea_font_name)
    if size:
        sz = etree.SubElement(rPr, qn('w:sz'))
        sz.set(qn('w:val'), str(int(size * 2)))
        szCs = etree.SubElement(rPr, qn('w:szCs'))
        szCs.set(qn('w:val'), str(int(size * 2)))
    if bold:
        etree.SubElement(rPr, qn('w:b'))
        etree.SubElement(rPr, qn('w:bCs'))
    t = etree.SubElement(r, qn('w:t'))
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    return new_p


def insert_after(doc, after_idx, text, bold=False, size=SZ_BODY, ea_font_name=FONT_EA):
    ref_element = doc.paragraphs[after_idx]._element
    new_p = make_para(doc, text, bold=bold, size=size, ea_font_name=ea_font_name)
    ref_element.addnext(new_p)


def find_text(doc, keyword, start=0):
    for i in range(start, len(doc.paragraphs)):
        if keyword in doc.paragraphs[i].text.strip():
            return i
    return -1


def count_chinese(text):
    return sum(1 for c in text if '一' <= c <= '鿿')


def main():
    doc = Document(OUT_PATH)
    initial_total = count_chinese(''.join(p.text.strip() for p in doc.paragraphs if p.text.strip()))
    print('Initial total Chinese:', initial_total)
    expansions = []

    # === Section 1.1 ===
    idx = find_text(doc, '显著提升大模型评价的效率')
    if idx >= 0:
        text = (
            '从评价方法的发展脉络来看，'
            '大模型评价经历了从人工评价'
            '到半自动化评价再到全自动化'
            '评价的演进过程。早期的大模'
            '型评价主要依赖人工构建的基'
            '准数据集，如MMLU、HellaSwag、HumanEval等，'
            '这些数据集在特定时期内发挥'
            '了重要作用，但其构建成本高、'
            '更新周期长，难以跟上大模型'
            '快速迭代的步伐。随后出现的'
            '评价框架如OpenAI Evals、EleutherAI LM Evaluation Harness等，'
            '提供了标准化的评价接口和预'
            '置数据集，但仍需人工维护数'
            '据集。近年来，基于智能体的'
            '自动评价方法开始兴起，通过'
            '将大模型自身的生成能力应用'
            '于评价数据集的构建过程，实'
            '现了评价流程的自动化。本课'
            '题正是在这一技术趋势下提出'
            '的，旨在探索一种更加灵活、'
            '自适应的自动评价智能体架构，'
            '能够在无需人工干预的情况下'
            '完成从数据集构建到模型评估'
            '的全流程。'
        )
        insert_after(doc, idx, text)
        expansions.append('1.1 research background')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 1.2 ===
    idx = find_text(doc, '在目前的研究中')
    if idx >= 0:
        text = (
            '在具体实施中，本研究采用了'
            '一种从点到面、逐步集成的研'
            '发策略。首先集中攻关评价数'
            '据集自动构建这一核心难题，'
            '设计并实现了基于多轮反馈自'
            '适应机制的题目生成智能体。'
            '该组件通过证据检索、文本分'
            '段、多策略采样和分阶段生成'
            '等技术手段，实现了从原始资'
            '料到高质量评价题目的全自动'
            '化转化。实验表明，该组件能'
            '够稳定产出覆盖多主题、多难'
            '度层级的题目集合，且生成过'
            '程具有良好的可控性和可复现'
            '性。随后，在数据集构建的基'
            '础上，进一步设计和实现了自'
            '动评价智能体架构，采用强中'
            '心化规划协同模式，将规划智'
            '能体作为全局决策核心，统一'
            '调度题目验证和模型评估等子'
            '任务。目前该架构已完成初版'
            '实现，能够支持从用户目标输'
            '入到多轮生成、验证、评估的'
            '完整闭环流程。对于第三部分'
            '的评估报告生成，已完成了数'
            '据结构设计和标准化输出格式'
            '的定义，可视化呈现和自然语'
            '言报告生成等后续工作正在规'
            '划中。'
        )
        insert_after(doc, idx, text)
        expansions.append('1.2 research progress')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.1(1): evidence retrieval ===
    idx = find_text(doc, '用户只需指定研究主题即可自动完成相关资料的检索和预处理')
    if idx >= 0:
        text = (
            '在检索策略方面，系统采用多'
            '轮递进式检索策略。首轮检索'
            '以用户指定的核心主题为查询'
            '词，从百科知识库中获取初步'
            '的候选文档列表；对于获取到'
            '的每篇文档，系统进一步提取'
            '其内部的关键概念和子主题作'
            '为扩展查询词，进行第二轮深'
            '度检索，以捕获与核心主题相'
            '关的周边知识和背景信息。这'
            '种递进式检索策略有效扩展了'
            '知识的覆盖范围，为后续的题'
            '目生成提供了丰富的素材基础'
            '。在文档摘要生成方面，系统'
            '采用了分层摘要策略：首先对'
            '每个文本块独立生成局部摘要'
            '，再综合所有局部摘要在文档'
            '层面生成全局摘要。这种分层'
            '策略既保留了局部细节信息，'
            '又提供了文档整体的主题概览'
            '，在证据采样和主题筛选阶段'
            '发挥着重要的指导作用。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.1(1) evidence retrieval')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.1(2): prompt design ===
    idx = find_text(doc, '初步筛选后的合格题目进入候选池')
    if idx >= 0:
        text = (
            '在题目生成的提示词设计方面'
            '，系统针对不同的题目模式和'
            '难度层级分别设计了专门化的'
            '提示词模板。对于问答题，提'
            '示词要求模型在生成问题时明'
            '确标注所需的推理类型，如事'
            '实性检索、逻辑推理、数值计'
            '算或综合分析等，并给出标准'
            '答案的同时标注答案在证据文'
            '本中的具体引用位置。对于多'
            '项选择题，提示词强调干扰项'
            '的设计原则：干扰项应当与正'
            '确答案在表面形式上保持一致'
            '，但在关键信息上与证据内容'
            '相矛盾。针对困难题目模式，'
            '提示词要求模型设计需要多步'
            '推理或多证据综合才能回答的'
            '问题，并展示完整的逻辑链条'
            '。此外，系统还支持通过示例'
            '引导的方式对生成格式进行约'
            '束，保证后续处理流程的兼容'
            '性和稳定性。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.1(2) prompt design')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.1(3): multi-round feedback signals ===
    idx = find_text(doc, '自主完成大规模、高质量评价数据集的构建')
    if idx >= 0:
        text = (
            '在多轮反馈的具体实现中，系'
            '统重点关注以下几类反馈信号'
            '的采集与应用。第一是产量类'
            '信号，记录每轮各模式和各难'
            '度层级的新增题目数量，用于'
            '评估当前生成策略的有效性；'
            '当某难度层级连续多轮产量低'
            '于预期时，系统会触发策略切'
            '换。第二是质量类信号，通过'
            '验证环节的通过率和质量评分'
            '，识别哪些主题或难度生成的'
            '题目质量偏低，从而在后续轮'
            '次中调整证据采样策略或切换'
            '提示词模板。第三是资源类信'
            '号，跟踪每个证据文本块的使'
            '用频率，当某个文本块被反复'
            '使用但产出质量持续偏低时，'
            '系统会降低其采样优先级。第'
            '四是收敛性信号，监控候选池'
            '总量的增速和新题目的边际增'
            '益，当增速显著放缓时，系统'
            '自动触发证据扩展流程，为目'
            '标主题补充新的资料，以维持'
            '生成效率。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.1(3) feedback signals')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.2(1): planner diagnosis ===
    idx = find_text(doc, '规划智能体还负责将上述计划组装为可执行的轮次规范')
    if idx >= 0:
        text = (
            '规划智能体的诊断阶段是多轮'
            '迭代的核心决策基础。在每一'
            '轮开始时，规划智能体综合前'
            '一轮各环节返回的执行结果数'
            '据，从多个维度进行状态评估'
            '：生成维度关注本轮各主题和'
            '难度层级的实际产量与计划目'
            '标的偏差程度、连续空轮情况'
            '以及生成失败率；验证维度关'
            '注验证各阶段的通过率和质量'
            '评分分布，判断当前验证策略'
            '的严格程度是否合适；评估维'
            '度关注各候选模型得分的分布'
            '特征，包括模型间的区分度水'
            '平和各题目的得分方差。基于'
            '这些诊断结果，规划智能体将'
            '其映射为策略调整指令，分别'
            '作用于后续的生成计划、控制'
            '参数配置和主题选择方向。这'
            '一诊断—规划—执行的闭环机'
            '制使得整个评价系统具备了动'
            '态自适应能力。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.2(1) diagnosis details')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.2(2): verify agent philosophy ===
    idx = find_text(doc, '实现验证策略的自适应优化')
    if idx >= 0:
        text = (
            '验证智能体的三阶段设计体现'
            '了从严格到灵活的递进式质量'
            '把关理念。引文验证阶段采用'
            '完全客观的规则化检查，逐条'
            '比对题目中的引用声明与证据'
            '原文，确保事实性信息的准确'
            '性和可追溯性。大模型质量评'
            '估阶段则引入了柔性评估维度'
            '，利用大模型自身的语义理解'
            '能力对题目的表述质量、答案'
            '合理性和逻辑一致性等进行综'
            '合判断。加权选择阶段将前两'
            '个阶段的评分结果进行综合，'
            '同时引入多样性和覆盖度约束'
            '，通过优化算法在有限的入选'
            '名额内选择出最具代表性的题'
            '目组合。这种刚柔并济的验证'
            '设计既保证了评价数据集的基'
            '本事实准确性，又兼顾了题目'
            '的质量和多样性，有效平衡了'
            '严格把关与效率之间的关系。'
            '此外，验证阶段产生的详细反'
            '馈记录为规划智能体的策略调'
            '优提供了重要的数据支撑，形'
            '成了生成—验证—反馈—调整'
            '的完整质量闭环。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.2(2) verify philosophy')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.2(3): model eval architecture ===
    idx = find_text(doc, '形成全链路的自适应优化')
    if idx >= 0:
        text = (
            '模型评估智能体在架构设计上'
            '采用了离线推理、异步评估的'
            '策略，将耗时较长的模型推理'
            '过程与计算相对轻量的指标评'
            '分过程分离，提升了整体评估'
            '流程的执行效率。在模型推理'
            '阶段，系统对多个候选模型同'
            '时发起推理请求，每个模型的'
            '调用通过统一的客户端接口进'
            '行抽象，支持不同厂商和不同'
            '部署方式的模型接入。推理结'
            '果以标准化格式缓存至本地存'
            '储，便于后续的离线分析。在'
            '指标评分阶段，系统支持多层'
            '次的评分体系：底层为自动评'
            '分指标，提供快速客观的基础'
            '评分；上层为大模型裁判评分'
            '，提供更深层的质性分析。两'
            '层次评分结果通过配置化的权'
            '重体系进行融合，形成每道题'
            '目的综合得分。最终的综合评'
            '估报告不仅包含各模型在各维'
            '度上的得分，还提供区分度分'
            '析、稳定性分析和一致性检验'
            '等高级分析功能。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.2(3) eval architecture')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.2(4): extensibility ===
    idx = find_text(doc, '具有良好的模块化程度和可扩展性')
    if idx >= 0:
        text = (
            '智能体之间的通信与协作通过'
            '标准化的数据接口实现。各子'
            '智能体之间的数据交换采用统'
            '一的数据结构规范，包括题目'
            '结构、验证结果结构、评估结'
            '果结构和反馈结构等。这种标'
            '准化设计确保了各组件之间的'
            '松耦合关系，使得任何一个子'
            '智能体的内部实现发生变化时'
            '，不会影响其他组件的正常运'
            '行。例如，验证智能体可以在'
            '不通知其他组件的情况下切换'
            '其内部的大模型质量评估策略'
            '，只要其输出的验证结果符合'
            '约定的数据结构，规划智能体'
            '和模型评估智能体即可无缝兼'
            '容。同样，模型评估智能体支'
            '持动态添加新的评估指标和裁'
            '判模型，新指标只需遵循约定'
            '的接口规范即可自动参与评估'
            '流程。这种可扩展的架构设计'
            '为本系统的持续演进提供了良'
            '好的基础。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.2(4) extensibility')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 2.3: report generation metrics ===
    idx = find_text(doc, '帮助研究人员快速理解各候选模型的能力差异')
    if idx >= 0:
        text = (
            '在评估报告的指标体系设计方'
            '面，系统支持多层次、可定制'
            '的评估指标配置。在基础指标'
            '层，系统计算每个候选模型在'
            '各题目上的得分以及在整体数'
            '据集上的汇总统计量，包括均'
            '值、中位数、标准差和最值分'
            '布等。在分析指标层，系统自'
            '动计算模型间区分度指标和数'
            '据集质量指标，包括区分度题'
            '目比例、过易题目比例、全失'
            '败题目比例等，帮助研究人员'
            '判断当前评价数据集和候选模'
            '型之间的匹配程度。在可视化'
            '方案方面，系统计划支持多种'
            '图表类型，包括雷达图、柱状'
            '图、热力图和散点图等。这些'
            '可视化图表与文本分析相结合'
            '，输出涵盖模型能力总览、分'
            '维度对比分析和典型题目案例'
            '等模块的综合报告，为研究人'
            '员提供直观、深入的模型能力'
            '分析体验。'
        )
        insert_after(doc, idx, text)
        expansions.append('2.3 metrics & visualization')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 3: schedule details ===
    idx = find_text(doc, '准备答辩材料，进行论文答辩')
    if idx >= 0:
        text = (
            '上述进度安排基于当前的研究'
            '基础和技术积累制定。由于系'
            '统的核心框架已经完成，各智'
            '能体组件已具备基本功能，后'
            '续工作主要围绕功能完善、系'
            '统集成和实验验证展开。评估'
            '报告生成模块的开发虽然尚在'
            '起步阶段，但其依赖的基础数'
            '据结构已经与模型评估智能体'
            '充分对接，技术方案明确，实'
            '现路径清晰。多轮反馈机制的'
            '优化工作将在已有实现的基础'
            '上进行增量改进，风险可控。'
            '系统集成测试和实验验证阶段'
            '将重点关注系统的稳定性、效'
            '率和可复现性。论文撰写将遵'
            '循硕士学位论文的规范要求，'
            '与实验验证工作同步推进。'
        )
        insert_after(doc, idx, text)
        expansions.append('3 schedule details')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 4: additional difficulty ===
    idx = find_text(doc, '第三，计算资源消耗问题')
    if idx >= 0:
        text = (
            '第四，评价结果的稳定性和可'
            '复现性问题。由于大模型生成'
            '过程具有一定的随机性，同一'
            '主题的多次生成可能得到不同'
            '的题目集合，导致评价结果在'
            '不同运行批次之间存在波动。'
            '虽然系统采用了随机种子控制'
            '和确定性采样策略来减小这种'
            '波动，但仍然难以完全消除。'
            '如何设计更加鲁棒的评估协议'
            '，使得评价结果对题目集合的'
            '随机波动不敏感，同时在有限'
            '的资源约束下获得统计显著的'
            '评估结论，是一个需要在理论'
            '和实践层面共同探索的问题。'
        )
        insert_after(doc, idx, text)
        expansions.append('4 difficulty 4')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # === Section 5: expand completion possibility ===
    idx = find_text(doc, '预计能够如期完成全部论文工作')
    if idx >= 0:
        text = (
            '从技术风险角度来看，系统的'
            '核心算法和架构已经过初步验'
            '证，不存在尚未攻克的关键技'
            '术难题。从实验条件来看，所'
            '需的大模型API服务和计算资源均'
            '已就位，具备开展实验的必要'
            '条件。从时间安排来看，当前'
            '距论文提交截止日期还有约11周'
            '时间，时间相对充裕，且已制'
            '定了详细的周计划，各项任务'
            '的时间估算合理。此外，导师'
            '在智能体和自然语言处理领域'
            '具有丰富的研究经验，能够在'
            '研究过程中提供及时的指导和'
            '建议。综上所述，按期完成全'
            '部论文工作具有较高的可行性'
            '。'
        )
        insert_after(doc, idx, text)
        expansions.append('5 completion details')
        doc.save(OUT_PATH)
        doc = Document(OUT_PATH)

    # ========== final count ==========
    doc = Document(OUT_PATH)
    all_text = ''.join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    chinese_total = count_chinese(all_text)

    # Per-section counts
    sections = {}
    current_sec = 'header'
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        if '课题研究的背景和意义' in txt:
            current_sec = '1.1'
        elif '课题研究内容和进展' in txt:
            current_sec = '1.2'
        elif '目前已经完成的研究工作' in txt:
            current_sec = '2'
        elif '后期拟完成的研究工作' in txt:
            current_sec = '3'
        elif '存在的困难与问题' in txt:
            current_sec = '4'
        elif '如期完成全部论文工作的可能性' in txt:
            current_sec = '5'
        elif txt == '参考文献':
            current_sec = 'ref'
        if current_sec not in sections:
            sections[current_sec] = ''
        sections[current_sec] += txt

    print('\n' + '=' * 40)
    print('Expansions performed:', len(expansions))
    for e in expansions:
        print('  +', e)
    print('\nPer-section Chinese chars:')
    for k in ['header', '1.1', '1.2', '2', '3', '4', '5', 'ref']:
        v = sections.get(k, '')
        c = count_chinese(v)
        print('  {}: {}'.format(k, c))
    print('\nTotal Chinese:', chinese_total)
    print('Added:', chinese_total - initial_total)
    print('=' * 40)

    # Check for code identifiers
    bad_patterns = ['GeneratorFeedback', 'ValidatorFeedback', 'EvaluatorFeedback',
                    'cold_start', 'gen_only', 'val_only']
    found_any = False
    for p in bad_patterns:
        if p.lower() in all_text.lower():
            print('WARNING: Found', p)
            found_any = True
    if not found_any:
        print('Code identifier check: PASSED')


if __name__ == '__main__':
    main()
