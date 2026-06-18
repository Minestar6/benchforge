from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


DOC_PATH = Path("/Users/zhaoziqing/Desktop/benchforge/report/赵子晴_中期报告_完善版.docx")
REPORT_PATH = Path(
    "/Users/zhaoziqing/Desktop/benchforge/runs/effectiveness_task_20260618_153413/effectiveness/task_effectiveness_report.json"
)


def set_run_style(run, *, bold: bool = False, size: float = 12.0) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)


def format_paragraph(paragraph, *, bold: bool = False, size: float = 12.0, first_line_cm: float = 0.74) -> None:
    paragraph.paragraph_format.first_line_indent = Cm(first_line_cm)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        set_run_style(run, bold=bold, size=size)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def move_table_after(paragraph: Paragraph, table) -> None:
    tbl = table._tbl
    tbl.getparent().remove(tbl)
    paragraph._p.addnext(tbl)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 11.0) -> None:
    cell.text = text
    for para in cell.paragraphs:
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.1
        for run in para.runs:
            set_run_style(run, bold=bold, size=size)


def load_metrics() -> dict:
    return json.loads(REPORT_PATH.read_text(encoding="utf-8"))


def ensure_not_inserted(doc: Document, marker: str) -> bool:
    return all(marker not in p.text for p in doc.paragraphs)


def main() -> None:
    metrics = load_metrics()
    doc = Document(DOC_PATH)

    marker = "2.1.3 多轮补充实验结果与分析"
    if not ensure_not_inserted(doc, marker):
        doc.save(DOC_PATH)
        return

    paragraphs = doc.paragraphs
    insert_after = None
    difficulty_para = None
    for para in paragraphs:
        txt = para.text.strip()
        if txt.startswith("在多轮反馈的具体实现中，系统重点关注以下几类反馈信号的采集与应用"):
            insert_after = para
        if txt.startswith("第四，评价结果的稳定性和可复现性问题"):
            difficulty_para = para
    if insert_after is None or difficulty_para is None:
        raise ValueError("Target insertion points not found in document.")

    b1 = metrics["b1"]
    bfull = metrics["bfull"]
    delta = metrics["benchmark_delta"]
    round1 = metrics["round_summaries"][0]
    round2 = metrics["round_summaries"][1]

    p = insert_paragraph_after(insert_after, marker)
    format_paragraph(p, bold=True, size=12.5, first_line_cm=0)

    p = insert_paragraph_after(
        p,
        (
            "为验证当前总框架在真实多轮场景下的运行效果，本研究在同一任务标识 "
            f"{metrics['task_id']} 下开展了两轮串行实验。两轮实验分别对应独立的运行目录 "
            "round_001 和 round_002，从而避免同轮数据相互覆盖；同时保持相同的主题设置（Artificial Intelligence、"
            "Quantum Computing）和统一的验证、评测流程，以便观察后续轮次对题目集合的增量贡献。题目生成与验证环节"
            "使用 DeepSeek-V4 系列配置，最终模型评测选择 deepseek-v4、deepseek-v4-pro 和 minimax 三个候选模型。"
        ),
    )
    format_paragraph(p)

    p = insert_paragraph_after(
        p,
        (
            f"从题目产出结果来看，第1轮最终入选 {round1['selected_count']} 道题，第2轮继续入选 {round2['selected_count']} 道题，"
            f"两轮累计得到 {metrics['all_rounds_selected_count']} 道高质量题目。其中，第1轮保留了 8 道问答题和 3 道选择题，"
            "第2轮保留了 10 道问答题和 3 道选择题，说明在相同任务目标下，第二轮能够继续为题库提供有效增量，而不是仅重复"
            "第一轮已经生成的题目。进一步观察主题分布可见，第1轮更偏向 Artificial Intelligence，第2轮则显著补充了 "
            "Quantum Computing 方向的题目，这与多轮生成在弱覆盖主题上继续扩展证据和题目的目标是一致的。"
        ),
    )
    format_paragraph(p)

    p = insert_paragraph_after(p, "表2.2 多轮补充实验结果汇总")
    format_paragraph(p, bold=True, size=12.0, first_line_cm=0)
    caption_para = p

    table = doc.add_table(rows=4, cols=7)
    if doc.tables:
        table._tbl.tblPr[:] = deepcopy(doc.tables[0]._tbl.tblPr[:])
    widths = [2.3, 1.6, 1.8, 1.8, 1.8, 1.9, 1.9]
    headers = ["题集", "题目数", "引用得分", "多样性得分", "Top-Bottom Gap", "得分方差", "区分题比例"]
    for idx, text in enumerate(headers):
        table.columns[idx].width = Cm(widths[idx])
        set_cell_text(table.rows[0].cells[idx], text, bold=True)

    rows = [
        [
            "第1轮入选题集(B1)",
            str(b1["question_count"]),
            f"{b1['evaluation_result']['dataset_metrics']['citation_score']['mean']:.4f}",
            f"{b1['evaluation_result']['dataset_metrics']['diversity_score']['score']:.4f}",
            f"{b1['benchmark_metrics']['top_bottom_gap']:.4f}",
            f"{b1['benchmark_metrics']['score_variance']:.6f}",
            f"{b1['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:.4f}",
        ],
        [
            "两轮累计题集(Bfull)",
            str(bfull["question_count"]),
            f"{bfull['evaluation_result']['dataset_metrics']['citation_score']['mean']:.4f}",
            f"{bfull['evaluation_result']['dataset_metrics']['diversity_score']['score']:.4f}",
            f"{bfull['benchmark_metrics']['top_bottom_gap']:.4f}",
            f"{bfull['benchmark_metrics']['score_variance']:.6f}",
            f"{bfull['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:.4f}",
        ],
        [
            "变化量(Bfull-B1)",
            f"+{metrics['later_round_selected_count']}",
            f"{bfull['evaluation_result']['dataset_metrics']['citation_score']['mean'] - b1['evaluation_result']['dataset_metrics']['citation_score']['mean']:+.4f}",
            f"{bfull['evaluation_result']['dataset_metrics']['diversity_score']['score'] - b1['evaluation_result']['dataset_metrics']['diversity_score']['score']:+.4f}",
            f"{delta['top_bottom_gap']:+.4f}",
            f"{delta['score_variance']:+.6f}",
            f"{bfull['evaluation_result']['discriminative_signals']['discriminative_question_ratio'] - b1['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:+.4f}",
        ],
    ]
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, text in enumerate(row):
            set_cell_text(table.rows[r_idx].cells[c_idx], text, size=10.5)
    p = insert_paragraph_after(
        caption_para,
        (
            "从表2.2可以看到，多轮补充实验首先验证了题库扩展能力：第二轮额外提供了13道入选题，使题目总数由11道扩展到24道。"
            f"同时，题集的平均引用得分由 {b1['evaluation_result']['dataset_metrics']['citation_score']['mean']:.4f} 提升到 "
            f"{bfull['evaluation_result']['dataset_metrics']['citation_score']['mean']:.4f}，多样性得分由 "
            f"{b1['evaluation_result']['dataset_metrics']['diversity_score']['score']:.4f} 提升到 "
            f"{bfull['evaluation_result']['dataset_metrics']['diversity_score']['score']:.4f}。这说明后续轮次并非仅机械地追加题目数量，"
            "而是在证据覆盖和语义分散度上为最终题集提供了额外增益，验证了多轮机制对题目集合质量的正向作用。"
        ),
    )
    format_paragraph(p)

    p = insert_paragraph_after(
        p,
        (
            "另一方面，从模型区分度指标来看，本次小样本实验并未呈现出单调增强的趋势。与仅使用第1轮入选题构成的题集相比，"
            f"两轮累计题集的 Top-Bottom Gap 从 {b1['benchmark_metrics']['top_bottom_gap']:.4f} 下降到 "
            f"{bfull['benchmark_metrics']['top_bottom_gap']:.4f}，得分方差也从 "
            f"{b1['benchmark_metrics']['score_variance']:.6f} 下降到 {bfull['benchmark_metrics']['score_variance']:.6f}。"
            f"不过，区分题比例由 {b1['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:.4f} "
            f"小幅上升到 {bfull['evaluation_result']['discriminative_signals']['discriminative_question_ratio']:.4f}，"
            "说明新增题目并未破坏题集的基本判别能力，但其增益更多体现在覆盖范围和多样性上，而不是直接放大模型间的总体分差。"
        ),
    )
    format_paragraph(p)

    p = insert_paragraph_after(
        p,
        (
            "综合来看，当前补充实验表明：本研究提出的多轮机制已经能够稳定带来题目数量扩展、引用质量提升和语义多样性增强；"
            "但在“如何让新增题目进一步强化模型区分度”这一点上，现阶段仍有优化空间。后续需要将区分度信号更显式地纳入反馈环节，"
            "例如根据模型间分差、题目方差或裁判分歧度来指导下一轮生成，从而使多轮反馈不仅提升题集规模和覆盖面，也更直接服务于最终评测效能。"
        ),
    )
    format_paragraph(p)

    move_table_after(caption_para, table)

    p = insert_paragraph_after(
        difficulty_para,
        (
            "最新的两轮补充实验进一步印证了这一问题：后续轮次能够显著提升题集规模、引用得分和多样性得分，但模型间总体分差"
            "（如Top-Bottom Gap和总体得分方差）并未同步提升。这说明当前反馈机制虽然能够有效补充题目覆盖，但对“高区分度题目”的"
            "定向挖掘还不够充分。后续将考虑把模型分差、题目方差和裁判不一致度等区分度信号显式纳入反馈与规划模块，使多轮生成的优化目标"
            "从“仅提高题目质量”进一步扩展为“同时提高评测区分能力”。"
        ),
    )
    format_paragraph(p)

    doc.save(DOC_PATH)


if __name__ == "__main__":
    main()
